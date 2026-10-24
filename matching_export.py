"""Formatted, safe Excel workbook for all Salla reconciliation tabs.

This module is intended to run inside the existing Streamlit project, where
openpyxl is already a declared dependency. No network writes occur here.
"""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO
import math
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEETS = (
    ("منتجات جديدة", "new"), ("اختلاف الأسماء", "names"),
    ("اختلاف حجم وعدد", "sizes"), ("اختلاف الأسعار", "prices"),
    ("اختلاف الضريبة", "taxes"), ("مشكلات SKU", "problems"),
    ("في سلة فقط", "salla_only"), ("استبعاد الأسماء", "name_exclusions"),
    ("استبعاد الأسعار", "price_exclusions"),
)
NAVY = "12243A"
TEAL = "077F80"
WHITE = "FFFFFF"
PALE = "E9F6F5"
ORANGE = "FFF0D5"
RED = "FDE9E9"


def _safe_xlsx_value(item):
    if item is None:
        return None
    if isinstance(item, Decimal):
        return float(item)
    try:
        if pd.isna(item):
            return None
    except (ValueError, TypeError):
        pass
    if isinstance(item, str):
        # Prevent spreadsheet formula injection from untrusted store item names.
        return "'" + item if re.match(r"^\s*[=+\-@]", item) else item
    if hasattr(item, "item"):
        try:
            item = item.item()
        except (ValueError, TypeError, AttributeError):
            pass
    if isinstance(item, (float, int, bool)) and (not isinstance(item, float) or math.isfinite(item)):
        return item
    return str(item)


def _style_table(sheet, fields, header_row=1):
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = f"B{header_row + 1}"
    sheet.row_dimensions[header_row].height = 33
    for cell in sheet[header_row]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", bold=True, color=WHITE, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = {
        "SKU": 20, "الاسم في النظام": 53, "الاسم في سلة": 53, "اسم المنتج": 53,
        "نوع اختلاف الاسم": 34, "الكمية/الحجم في النظام": 27, "الكمية/الحجم في سلة": 27,
        "درجة تأكيد المقاس": 22, "المشكلة": 33, "سبب الاستبعاد": 35,
    }
    for idx, field in enumerate(fields, 1):
        sheet.column_dimensions[get_column_letter(idx)].width = widths.get(field, min(29, max(17, len(str(field)) * 1.3)))
    if sheet.max_row > header_row:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(fields))}{sheet.max_row}"


def _fill_sheet(sheet, frame: pd.DataFrame, highlight_size=False):
    fields = [str(x) for x in frame.columns]
    if not fields:
        fields = ["البيان"]
        frame = pd.DataFrame(columns=fields)
    sheet.append(fields)
    for index, record in enumerate(frame.itertuples(index=False, name=None), 2):
        sheet.append([_safe_xlsx_value(value) for value in record])
        if highlight_size and "نوع اختلاف الاسم" in fields:
            kind = str(record[fields.index("نوع اختلاف الاسم")])
            if kind != "اختلاف نص الاسم":
                fill = PatternFill("solid", fgColor=RED if "تركيز/جرعة" in kind else
                                   ORANGE if "اختلاف " in kind else "FFF7DB")
                for cell in sheet[index]:
                    cell.fill = fill
        if index % 2 == 0 and not highlight_size:
            for cell in sheet[index]:
                cell.fill = PatternFill("solid", fgColor="F4F8FC")
    _style_table(sheet, fields)
    # The SKU is text; preserving leading zeros is fundamental for matching.
    if "SKU" in fields:
        col = fields.index("SKU") + 1
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=col, max_col=col):
            for cell in row:
                cell.number_format = "@"


def build_full_reconciliation_excel(report: dict, price_tolerance="0.01") -> bytes:
    wb = Workbook()
    overview = wb.active
    overview.title = "ملخص"
    overview.sheet_view.rightToLeft = True
    overview.append(["تقرير مطابقة منتجات سلة — الأسعار قبل الضريبة", "القيمة"])
    lines = [
        ("أساس المقارنة", "السعر في كلا الملفين غير شامل الضريبة"),
        ("هامش الأسعار قبل الضريبة", str(price_tolerance)),
        ("ملاحظة", "التبويبات غير حصرية؛ اختلاف المنتج قد يظهر في أكثر من تبويب"),
        ("سلامة البيانات", "لا تُجرى أي تعديلات على سلة من هذا التقرير"),
    ]
    titles = {
        "salla_rows": "صفوف سلة", "system_rows": "صفوف النظام",
        "comparable_common_skus": "SKU مشتركة", "new_products": "منتجات جديدة",
        "ready_to_upload": "صالحة لإعداد ملف الرفع", "salla_only": "سلة فقط",
        "different_names": "اختلافات أسماء (بعد الاستبعاد)", "size_count_differences": "فروقات العبوة: مؤكدة ومحتملة", "confirmed_size_count_differences": "اختلاف مقاس مؤكد في الوصف",
        "different_prices": "اختلافات أسعار (بعد الاستبعاد)", "different_taxes": "اختلافات الضريبة",
        "problems": "مشكلات البيانات", "ambiguous_skus": "SKU ملتبسة",
        "name_exclusions_applied": "SKU مستبعدة من مقارنة الأسماء", "price_exclusions_applied": "SKU مستبعدة من مقارنة الأسعار",
        "excluded_system_rows": "صفوف تصنيفات مستبعدة",
    }
    lines.extend((titles.get(k, k), v) for k, v in report["summary"].items())
    for field, value in lines:
        overview.append([field, value])
    _style_table(overview, ["البيان", "القيمة"])
    overview.column_dimensions["A"].width = 48
    overview.column_dimensions["B"].width = 65
    for title, key in SHEETS:
        ws = wb.create_sheet(title)
        _fill_sheet(ws, report[key], highlight_size=key in ("names", "sizes"))
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def build_exclusion_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "SKU للاستبعاد"
    ws.append(["SKU", "ملاحظة اختيارية"])
    # Intentionally no example SKU: prevent accidental exclusion of a real product.
    _style_table(ws, ["SKU", "ملاحظة اختيارية"])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 62
    out = BytesIO(); wb.save(out)
    return out.getvalue()
