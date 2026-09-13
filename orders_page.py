import streamlit as st
import pandas as pd
import io
import time
from datetime import datetime, timedelta
import concurrent.futures
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from utils import get_headers, safe_api_request, SALLA_API_URL

def get_orders_list(from_date, to_date, headers, search_keyword=None, order_refs_list=None):
    """سحب الطلبات: يدعم التواريخ، كلمة بحث، أو قائمة أرقام طلبات من ملف"""
    orders = []
    status_text = st.empty()
    progress_bar = st.progress(0)
    
    if order_refs_list and len(order_refs_list) > 0:
        total = len(order_refs_list)
        for idx, ref in enumerate(order_refs_list):
            status_text.info(f"📥 جاري سحب الطلب رقم {ref} ({idx+1} من {total})...")
            url = f"https://api.salla.dev/admin/v2/orders?keyword={ref}"
            res = safe_api_request("GET", url, headers)
            if res and res.get("data"):
                matched = [o for o in res["data"] if str(o.get('reference_id')) == str(ref)]
                orders.extend(matched if matched else res["data"])
            progress_bar.progress((idx + 1) / total)
    else:
        page = 1
        total_pages = 1
        while page <= total_pages:
            status_text.info(f"📥 جاري حصر الطلبات (صفحة {page} من {total_pages if page > 1 else '...'})...")
            url = f"https://api.salla.dev/admin/v2/orders?from_date={from_date}&to_date={to_date}&per_page=50&page={page}"
            if search_keyword: url += f"&keyword={search_keyword}"
                
            res = safe_api_request("GET", url, headers)
            if not res or not res.get("data"): break
            if page == 1: total_pages = res.get("pagination", {}).get("totalPages", 1)
            
            orders.extend(res["data"])
            progress_bar.progress(min(page / total_pages, 1.0))
            page += 1
            
    progress_bar.empty()
    status_text.empty()
    return orders

def fetch_single_order_details(order_id, headers):
    """دالة مساعدة لسحب تفاصيل طلب واحد (تستخدم للتنفيذ المتوازي السريع)"""
    res_order = safe_api_request("GET", f"https://api.salla.dev/admin/v2/orders/{order_id}", headers)
    order_data = res_order.get("data", {}) if res_order else {}
    
    items_res = safe_api_request("GET", f"https://api.salla.dev/admin/v2/orders/items?order_id={order_id}", headers)
    order_items = items_res.get("data", []) if items_res else []
    
    items_data = order_items.copy() if isinstance(order_items, list) else []

    for inv_item in items_data:
        pid = inv_item.get('product_id') or (inv_item.get('product', {}).get('id'))
        if pid and str(pid).isdigit():
            prod_info = next((p for p in st.session_state.get('all_products', []) if str(p.get('id')) == str(pid)), None)
            if not prod_info:
                p_res = safe_api_request("GET", f"https://api.salla.dev/admin/v2/products/{pid}", headers)
                prod_info = p_res.get("data", {}) if p_res else {}
                
            if prod_info and prod_info.get('type') == 'group_products':
                if prod_info.get('grouped_items'):
                    inv_item['grouped_items'] = prod_info.get('grouped_items')
                elif prod_info.get('consisted_products'):
                    inv_item['consisted_products'] = prod_info.get('consisted_products')
                    
    order_data['items'] = items_data
    return order_data

def get_detailed_orders(orders_summary, headers):
    """⚡ سحب التفاصيل الدقيقة بأسلوب متوازي فائق السرعة (Multithreading)"""
    detailed_orders = []
    status_text = st.empty()
    progress_bar = st.progress(0)
    total = len(orders_summary)
    completed = 0
    
    status_text.info(f"⚡ جاري سحب تفاصيل {total} طلب بوضع السرعة القصوى (Parallel Processing)...")
    
    with concurrent.futures.ThreadPoolExecutor(max_workers=8) as executor:
        future_to_order = {
            executor.submit(fetch_single_order_details, o_sum.get("id"), headers): o_sum 
            for o_sum in orders_summary
        }
        
        for future in concurrent.futures.as_completed(future_to_order):
            completed += 1
            progress_bar.progress(completed / total)
            try:
                data = future.result()
                if data: detailed_orders.append(data)
            except Exception as exc:
                print(f"Order fetch generated an exception: {exc}")
                
    status_text.success(f"✅ تم سحب تفاصيل {len(detailed_orders)} طلب بنجاح وبسرعة فائقة!")
    progress_bar.empty()
    return detailed_orders

def extract_shipping_company(order):
    """دالة مطابقة تماماً لواجهة سلة لاستخراج اسم شركة الشحن الحقيقية"""
    # 1. فحص كائن الـ shipping المباشر (مثل شركة بيز وغيرها)
    shipping_obj = order.get('shipping', {})
    if isinstance(shipping_obj, dict):
        comp = shipping_obj.get('company')
        if comp and comp != 'غير متوفر': return comp
        
    # 2. فحص مصفوفة الشحنات إن وجدت
    shipments = order.get('shipments', [])
    if shipments and isinstance(shipments, list):
        c_name = shipments[0].get('courier_name')
        if c_name: return c_name
        
    return "شحن يدوي / عادي"

def process_financials(orders):
    """أداة المعالجة المحاسبية: تفصيل المنتجات وإضافة سطر مستقل للشحن"""
    detailed_rows = []
    taxable_stats = {'item_sales': 0.0, 'shipping_sales': 0.0, 'qty': 0, 'item_tax': 0.0, 'shipping_tax': 0.0}
    nontaxable_stats = {'item_sales': 0.0, 'shipping_sales': 0.0, 'qty': 0, 'item_tax': 0.0, 'shipping_tax': 0.0}
    
    for order in orders:
        subtotal = float(order.get('amounts', {}).get('sub_total', {}).get('amount', 0))
        shipping_cost = float(order.get('amounts', {}).get('shipping_cost', {}).get('amount', 0))
        
        tax_obj = order.get('amounts', {}).get('tax', {})
        if isinstance(tax_obj, dict) and isinstance(tax_obj.get('amount'), dict):
            order_total_tax = float(tax_obj['amount'].get('amount', 0))
        else:
            order_total_tax = 0.0
            
        discounts = order.get('amounts', {}).get('discounts', [])
        total_coupon = sum(float(d.get('discount', 0)) for d in discounts if d.get('type') != 'special_offer' and 'عرض' not in str(d.get('title', '')))
        all_special_offers = [d for d in discounts if d.get('type') == 'special_offer' or 'عرض' in str(d.get('title', ''))]
        
        order_skus = [str(item.get('sku', '')).strip() for item in order.get('items', [])]
        specific_offers = []
        general_offers_total = 0.0
        
        for sp in all_special_offers:
            sp_title = str(sp.get('title', ''))
            matched = any(sku and sku in sp_title for sku in order_skus)
            if matched: specific_offers.append(sp)
            else: general_offers_total += float(sp.get('discount', 0))

        order_branches = order.get('order_branches', [])
        branch = order_branches[0].get('name', 'الفرع الرئيسي') if order_branches else 'غير متوفر'
        shipping_company = extract_shipping_company(order)

        order_items_tax_sum = 0.0 
        items = order.get('items', [])
        
        for item in items:
            sku = str(item.get('sku', 'غير متوفر')).strip()
            main_qty = int(item.get('quantity', 1))
            
            price_without_tax = 0.0
            if item.get('amounts') and item['amounts'].get('price_without_tax'):
                price_without_tax = float(item['amounts']['price_without_tax'].get('amount', 0))
            elif item.get('price'):
                price_without_tax = float(item['price'].get('amount', 0))
                
            tax_node = item.get('amounts', {}).get('tax', {}) if 'amounts' in item else item.get('tax', {})
            tax_percent_raw = tax_node.get('percent', 15.0)
            tax_percent = float(tax_percent_raw) if tax_percent_raw is not None else 15.0
            
            item_subtotal = price_without_tax * main_qty
            
            ratio = (item_subtotal / subtotal) if subtotal > 0 else 0
            item_coupon_share = ratio * total_coupon
            item_general_special_share = ratio * general_offers_total
            
            item_specific_special_share = sum(float(sp.get('discount', 0)) for sp in specific_offers if sku != 'غير متوفر' and sku in str(sp.get('title', '')))
            total_special_share = item_general_special_share + item_specific_special_share
            
            item_total_after_disc = item_subtotal - item_coupon_share - total_special_share
            if item_total_after_disc < 0: item_total_after_disc = 0
            
            is_taxable = tax_percent > 0
            calculated_tax = item_total_after_disc * (tax_percent / 100) if is_taxable else 0.0
            order_items_tax_sum += calculated_tax
            
            item_net_sales = item_total_after_disc + calculated_tax 
            
            sub_items_extracted = []
            if item.get('grouped_items'):
                for gi in item.get('grouped_items'):
                    sub_prod = gi.get('product', {}) if isinstance(gi.get('product'), dict) else {}
                    sub_name = sub_prod.get('name', 'بدون اسم')
                    sub_sku = sub_prod.get('sku', '')
                    sub_qty = int(gi.get('quantity', 1)) * main_qty
                    sub_label = f"{sub_name} (SKU: {sub_sku})" if sub_sku else sub_name
                    sub_items_extracted.append({"name": sub_label, "qty": sub_qty})
            elif item.get('consisted_products'):
                for cp in item.get('consisted_products'):
                    sub_name = cp.get('name', 'بدون اسم')
                    sub_sku = cp.get('sku', '')
                    sub_qty = int(cp.get('quantity_in_group', 1)) * main_qty
                    sub_label = f"{sub_name} (SKU: {sub_sku})" if sub_sku else sub_name
                    sub_items_extracted.append({"name": sub_label, "qty": sub_qty})
            
            for opt in item.get('options', []):
                opt_name = opt.get('name', '')
                val = opt.get('value')
                if isinstance(val, list):
                    for v in val: sub_items_extracted.append({"name": f"{opt_name}: {v.get('name', '')}", "qty": main_qty})
                elif isinstance(val, dict):
                    sub_items_extracted.append({"name": f"{opt_name}: {val.get('name', '')}", "qty": main_qty})
                else:
                    sub_items_extracted.append({"name": f"{opt_name}: {val}", "qty": main_qty})
            
            if not sub_items_extracted:
                sub_items_extracted = [{"name": "بدون", "qty": main_qty}]
                
            qty_for_stats = sum(si['qty'] for si in sub_items_extracted) if sub_items_extracted[0]['name'] != "بدون" else main_qty
            
            if is_taxable:
                taxable_stats['item_sales'] += item_total_after_disc
                taxable_stats['qty'] += qty_for_stats
                taxable_stats['item_tax'] += calculated_tax
            else:
                nontaxable_stats['item_sales'] += item_total_after_disc
                nontaxable_stats['qty'] += qty_for_stats
            
            splits = len(sub_items_extracted)
            for sub in sub_items_extracted:
                detailed_rows.append({
                    "الفرع": branch,
                    "تاريخ الطلب": str(order.get('date', {}).get('date', ''))[:10],
                    "رقم الطلب": order.get('reference_id', ''),
                    "حالة الطلب": order.get('status', {}).get('name', ''),
                    "اسم العميل": f"{order.get('customer', {}).get('first_name', '')} {order.get('customer', {}).get('last_name', '')}".strip(),
                    "المدينة": order.get('customer', {}).get('city', ''),
                    "شركة الشحن": shipping_company,
                    "رقم الصنف (SKU)": sku,
                    "اسم الصنف": item.get('name', ''),
                    "الأصناف الفرعية": sub['name'],
                    "خاضع للضريبة": "نعم" if is_taxable else "لا",
                    "الكمية": sub['qty'],
                    "سعر الصنف (بدون ضريبة)": price_without_tax / splits,
                    "قيمة خصم الكوبون": round(item_coupon_share / splits, 2),
                    "قيمة خصم العرض الخاص": round(total_special_share / splits, 2),
                    "الاجمالي بعد الخصم": round(item_total_after_disc / splits, 2),
                    "تكلفة الشحن": 0.0,
                    "الضريبة": round(calculated_tax / splits, 2),
                    "صافي المبيعات": round(item_net_sales / splits, 2)
                })
        
        shipping_tax = max(0.0, order_total_tax - order_items_tax_sum)
        shipping_is_taxable = shipping_tax > 0 or order_total_tax > 0
        
        if shipping_cost > 0:
            if shipping_is_taxable:
                taxable_stats['shipping_sales'] += shipping_cost
                taxable_stats['shipping_tax'] += shipping_tax
            else:
                nontaxable_stats['shipping_sales'] += shipping_cost
                
            detailed_rows.append({
                "الفرع": branch,
                "تاريخ الطلب": str(order.get('date', {}).get('date', ''))[:10],
                "رقم الطلب": order.get('reference_id', ''),
                "حالة الطلب": order.get('status', {}).get('name', ''),
                "اسم العميل": f"{order.get('customer', {}).get('first_name', '')} {order.get('customer', {}).get('last_name', '')}".strip(),
                "المدينة": order.get('customer', {}).get('city', ''),
                "شركة الشحن": shipping_company,
                "رقم الصنف (SKU)": "SHIPPING",
                "اسم الصنف": "تكلفة الشحن",
                "الأصناف الفرعية": "بدون",
                "خاضع للضريبة": "نعم" if shipping_is_taxable else "لا",
                "الكمية": 1,
                "سعر الصنف (بدون ضريبة)": shipping_cost,
                "قيمة خصم الكوبون": 0.0,
                "قيمة خصم العرض الخاص": 0.0,
                "الاجمالي بعد الخصم": shipping_cost,
                "تكلفة الشحن": shipping_cost,
                "الضريبة": round(shipping_tax, 2),
                "صافي المبيعات": round(shipping_cost + shipping_tax, 2)
            })
                
    return detailed_rows, taxable_stats, nontaxable_stats

def generate_short_export(orders):
    """بناء التصدير المختصر"""
    rows = []
    columns_list = [
        "الفرع", "تاريخ الطلب", "رقم الطلب", "حالة الطلب", "اسم العميل",
        "رقم الجوال", "بريد العميل", "المدينة", "شركة الشحن", "طريقة الدفع",
        "utm_source", "مجموع السلة", "الخصم الإجمالي", "قيمة خصم الكوبون",
        "قيمة خصم العروض الخاصة", "الاجمالي بعد الخصم", "تكلفة الشحن",
        "الضريبة", "صافي المبيعات", "المبلغ المسترجع"
    ]
    
    for order in orders:
        discounts = order.get('amounts', {}).get('discounts', [])
        coupon_disc = sum(float(d.get('discount', 0)) for d in discounts if d.get('type') != 'special_offer' and 'عرض' not in str(d.get('title', '')))
        special_disc = sum(float(d.get('discount', 0)) for d in discounts if d.get('type') == 'special_offer' or 'عرض' in str(d.get('title', '')))
        total_discount = coupon_disc + special_disc
        
        subtotal = float(order.get('amounts', {}).get('sub_total', {}).get('amount', 0))
        shipping = float(order.get('amounts', {}).get('shipping_cost', {}).get('amount', 0))
        tax = float(order.get('amounts', {}).get('tax', {}).get('amount', {}).get('amount', 0))
        refund = float(order.get('payment_actions', {}).get('refund_action', {}).get('refund_amount', {}).get('amount', 0))
        
        total_after_disc = subtotal - total_discount
        net_sales_with_tax = total_after_disc + tax + shipping 
        
        order_branches = order.get('order_branches', [])
        branch = order_branches[0].get('name', 'الفرع الرئيسي') if order_branches else 'غير متوفر'
        shipping_company = extract_shipping_company(order)

        rows.append({
            "الفرع": branch,
            "تاريخ الطلب": str(order.get('date', {}).get('date', ''))[:10],
            "رقم الطلب": order.get('reference_id', ''),
            "حالة الطلب": order.get('status', {}).get('name', ''),
            "اسم العميل": f"{order.get('customer', {}).get('first_name', '')} {order.get('customer', {}).get('last_name', '')}".strip(),
            "رقم الجوال": f"{order.get('customer', {}).get('mobile_code', '')}{order.get('customer', {}).get('mobile', '')}",
            "بريد العميل": order.get('customer', {}).get('email', ''),
            "المدينة": order.get('customer', {}).get('city', ''),
            "شركة الشحن": shipping_company,
            "طريقة الدفع": order.get('payment_method', ''),
            "utm_source": order.get('source_details', {}).get('utm_source', '') or order.get('campaign', {}).get('source', 'مباشر'),
            "مجموع السلة": subtotal,
            "الخصم الإجمالي": total_discount,
            "قيمة خصم الكوبون": coupon_disc,
            "قيمة خصم العروض الخاصة": special_disc,
            "الاجمالي بعد الخصم": total_after_disc,
            "تكلفة الشحن": shipping,
            "الضريبة": tax,
            "صافي المبيعات": net_sales_with_tax,
            "المبلغ المسترجع": refund
        })
        
    df = pd.DataFrame(rows, columns=columns_list)
    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "التصدير المختصر"
    ws.sheet_view.rightToLeft = True
    
    headers = list(df.columns)
    ws.append(headers)
    for row in df.itertuples(index=False, name=None): ws.append(row)
        
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="00EBCF", bold=True, size=12)
    center_align = Alignment(horizontal="center", vertical="center")
    
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill; cell.font = header_font; cell.alignment = center_align
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18
        
    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}{ws.max_row}"
    wb.save(buf)
    return buf.getvalue()

def generate_detailed_export(detailed_rows, taxable_stats, nontaxable_stats):
    """بناء إكسيل التصدير التفصيلي"""
    columns_list = [
        "الفرع", "تاريخ الطلب", "رقم الطلب", "حالة الطلب", "اسم العميل",
        "المدينة", "شركة الشحن", "رقم الصنف (SKU)", "اسم الصنف", "الأصناف الفرعية", "خاضع للضريبة",
        "الكمية", "سعر الصنف (بدون ضريبة)", "قيمة خصم الكوبون", "قيمة خصم العرض الخاص",
        "الاجمالي بعد الخصم", "تكلفة الشحن", "الضريبة", "صافي المبيعات"
    ]
    df = pd.DataFrame(detailed_rows, columns=columns_list)
    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "التصدير التفصيلي"
    ws.sheet_view.rightToLeft = True
    
    ws.merge_cells('A1:G1')
    ws['A1'] = "📊 إحصائيات المبيعات التفصيلية (شامل الشحن والضريبة)"
    ws['A1'].font = Font(bold=True, size=13, color="FFFFFF")
    ws['A1'].fill = PatternFill(start_color="8E44AD", end_color="8E44AD", fill_type="solid")
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
    
    stat_headers = ["النوع", "مبيعات المنتجات", "مبيعات الشحن", "إجمالي المبيعات", "الكمية المباعة", "ضريبة المنتجات", "ضريبة الشحن", "إجمالي الضريبة"]
    ws.append(stat_headers)
    for cell in ws[2]: cell.font = Font(bold=True); cell.fill = PatternFill(start_color="ECF0F1", fill_type="solid")
        
    ws.append([
        "خاضعة للضريبة", round(taxable_stats['item_sales'], 2), round(taxable_stats['shipping_sales'], 2), 
        round(taxable_stats['item_sales'] + taxable_stats['shipping_sales'], 2), taxable_stats['qty'],
        round(taxable_stats['item_tax'], 2), round(taxable_stats['shipping_tax'], 2), 
        round(taxable_stats['item_tax'] + taxable_stats['shipping_tax'], 2)
    ])
    ws.append([
        "غير خاضعة للضريبة", round(nontaxable_stats['item_sales'], 2), round(nontaxable_stats['shipping_sales'], 2), 
        round(nontaxable_stats['item_sales'] + nontaxable_stats['shipping_sales'], 2), nontaxable_stats['qty'],
        0.0, 0.0, 0.0
    ])
    ws.append([])
    
    headers = list(df.columns)
    ws.append(headers)
    header_row_idx = ws.max_row
    for row in df.itertuples(index=False, name=None): ws.append(row)
        
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="00EBCF", bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col)
        cell.fill = header_fill; cell.font = header_font; cell.alignment = center_align
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18
        
    ws.auto_filter.ref = f"A{header_row_idx}:{openpyxl.utils.get_column_letter(len(headers))}{ws.max_row}"
    wb.save(buf)
    return buf.getvalue()

def render_orders_page():
    st.markdown("""
    <div style="background: linear-gradient(135deg, #0F1C2E 0%, #8E44AD 100%); padding: 15px 25px; border-radius: 12px; color: white; margin-bottom: 20px; box-shadow: 0 4px 6px rgba(0,0,0,0.1);">
        <h2 style="color: white; margin: 0;">📦 تفاصيل طلبات المتجر (التصدير المحاسبي)</h2>
    </div>
    """, unsafe_allow_html=True)
    
    headers = get_headers()
    if not headers: return
    
    with st.container(border=True):
        st.markdown("#### 🔍 أدوات البحث واستخراج الطلبات")
        
        uploaded_orders_file = st.file_uploader("📂 رفع ملف أرقام الطلبات (اختياري - Excel/CSV):", type=["xlsx", "csv"], help="ضع أرقام الطلبات في العمود الأول للملف.")
        order_refs_from_file = []
        if uploaded_orders_file:
            try:
                df_refs = pd.read_excel(uploaded_orders_file) if uploaded_orders_file.name.endswith('.xlsx') else pd.read_csv(uploaded_orders_file)
                order_refs_from_file = df_refs.iloc[:, 0].dropna().astype(str).tolist()
                st.success(f"✅ تم قراءة {len(order_refs_from_file)} رقم طلب من الملف.")
            except Exception as e:
                st.error("خطأ في قراءة الملف.")
        
        search_query = st.text_input("🔎 ابحث برقم الطلب أو الجوال (اختياري):", placeholder="مثال: 41027662", disabled=bool(order_refs_from_file))
        
        col1, col2, col3 = st.columns([2, 2, 1])
        with col1: from_date = st.date_input("من تاريخ:", value=datetime.now().date() - timedelta(days=7), disabled=bool(order_refs_from_file))
        with col2: to_date = st.date_input("إلى تاريخ:", value=datetime.now().date(), disabled=bool(order_refs_from_file))
        with col3:
            st.markdown("<br>", unsafe_allow_html=True)
            if st.button("🚀 سحب الطلبات", use_container_width=True, type="primary"):
                with st.spinner("جاري السحب بالسرعة القصوى..."):
                    orders_summary = get_orders_list(
                        from_date.strftime('%Y-%m-%d'), to_date.strftime('%Y-%m-%d'), headers,
                        search_keyword=search_query.strip() if search_query else None,
                        order_refs_list=order_refs_from_file
                    )
                if not orders_summary:
                    st.warning("⚠️ لا توجد طلبات مطابقة.")
                    st.session_state['detailed_fetched_orders'] = []
                else:
                    detailed_orders = get_detailed_orders(orders_summary, headers)
                    st.session_state['detailed_fetched_orders'] = detailed_orders

    if st.session_state.get('detailed_fetched_orders'):
        orders_data = st.session_state['detailed_fetched_orders']
        
        detailed_rows, t_stats, nt_stats = process_financials(orders_data)
        
        st.markdown("---")
        st.markdown("### 📊 إحصائيات المبيعات التفصيلية")
        
        t_total_sales = t_stats['item_sales'] + t_stats['shipping_sales']
        t_total_tax = t_stats['item_tax'] + t_stats['shipping_tax']
        nt_total_sales = nt_stats['item_sales'] + nt_stats['shipping_sales']
        
        stats_html = f"""
<div style="display:flex; gap:15px; margin-bottom: 25px; flex-wrap: wrap;">
<div style="flex:1; min-width: 300px; background:linear-gradient(135deg, #16a085, #1abc9c); padding:20px; border-radius:12px; color:white; text-align:center; box-shadow: 0 4px 10px rgba(0,0,0,0.15);">
<h4 style="margin:0; font-size:16px; color:#e0f7fa;">✅ مبيعات خاضعة للضريبة</h4>
<h2 style="margin:10px 0; font-size:28px;">{round(t_total_sales, 2):,} <span style="font-size:16px;">SAR</span></h2>
<div style="font-size:13px; background:rgba(0,0,0,0.15); padding:8px; border-radius:8px; margin-bottom:8px;">
<b>تفصيل المبيعات:</b> المنتجات: {round(t_stats['item_sales'], 2):,} | الشحن: {round(t_stats['shipping_sales'], 2):,}
</div>
<div style="font-size:13px; background:rgba(0,0,0,0.15); padding:8px; border-radius:8px;">
<b>إجمالي الضريبة:</b> {round(t_total_tax, 2):,} SAR <br> (المنتجات: {round(t_stats['item_tax'], 2):,} | الشحن: {round(t_stats['shipping_tax'], 2):,})
</div>
</div>
<div style="flex:1; min-width: 300px; background:linear-gradient(135deg, #34495e, #2c3e50); padding:20px; border-radius:12px; color:white; text-align:center; box-shadow: 0 4px 10px rgba(0,0,0,0.15);">
<h4 style="margin:0; font-size:16px; color:#ecf0f1;">🚫 مبيعات غير خاضعة للضريبة</h4>
<h2 style="margin:10px 0; font-size:28px;">{round(nt_total_sales, 2):,} <span style="font-size:16px;">SAR</span></h2>
<div style="font-size:13px; background:rgba(0,0,0,0.15); padding:8px; border-radius:8px; margin-bottom:8px;">
<b>تفصيل المبيعات:</b> المنتجات: {round(nt_stats['item_sales'], 2):,} | الشحن: {round(nt_stats['shipping_sales'], 2):,}
</div>
<div style="font-size:13px; background:rgba(0,0,0,0.15); padding:8px; border-radius:8px;">
<b>إجمالي الكمية المباعة:</b> {nt_stats['qty']} وحدة
</div>
</div>
</div>
"""
        st.markdown(stats_html, unsafe_allow_html=True)
                
        st.markdown("---")
        st.markdown("### 📥 خيارات التصدير")
        
        col_short, col_detailed = st.columns(2)
        with col_short:
            excel_short = generate_short_export(orders_data)
            st.download_button(label="📥 تحميل تصدير إكسيل المختصر", data=excel_short, file_name=f"Orders_Short_{datetime.now().strftime('%Y%m%d')}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True, type="primary")
            
        with col_detailed:
            excel_detailed = generate_detailed_export(detailed_rows, t_stats, nt_stats)
            st.download_button(label="📥 تحميل تصدير إكسيل التفصيلي (شامل الفرعيات)", data=excel_detailed, file_name=f"Orders_Detailed_{datetime.now().strftime('%Y%m%d')}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True, type="primary")
            
        st.markdown("---")
        st.markdown(f"### 📋 ملخص الطلبات المسحوبة ({len(orders_data)})")
        
        cols = st.columns(2)
        for i, o in enumerate(orders_data):
            col = cols[i % 2]
            
            c_name = f"{o.get('customer', {}).get('first_name', '')} {o.get('customer', {}).get('last_name', '')}".strip()
            city = o.get('customer', {}).get('city', 'غير محدد')
            o_date = str(o.get('date', {}).get('date', ''))[:16]
            status_name = o.get('status', {}).get('name', 'غير محدد')
            
            shipping_company = extract_shipping_company(o)
            order_branches = o.get('order_branches', [])
            branch = order_branches[0].get('name', 'الفرع الرئيسي') if order_branches else 'غير متوفر'
            utm_source = o.get('source_details', {}).get('utm_source', '') or o.get('campaign', {}).get('source', 'مباشر')
            
            subtotal = float(o.get('amounts', {}).get('sub_total', {}).get('amount', 0))
            tax = float(o.get('amounts', {}).get('tax', {}).get('amount', {}).get('amount', 0))
            shipping_cost = float(o.get('amounts', {}).get('shipping_cost', {}).get('amount', 0))
            discounts = o.get('amounts', {}).get('discounts', [])
            total_discount = sum(float(d.get('discount', 0)) for d in discounts)
            o_total = float(o.get('amounts', {}).get('total', {}).get('amount', 0))
            
            border_color = "#2ecc71" if "تنفيذ" in status_name or "توصيل" in status_name else ("#e74c3c" if "لغي" in status_name else "#f39c12")
            
            with col:
                card_html = f"""
<div style="background: linear-gradient(145deg, #1e293b, #0f172a); border-radius: 12px; padding: 16px; margin-bottom: 16px; border: 1px solid #334155; border-right: 5px solid {border_color}; position: relative; box-shadow: 0 4px 15px rgba(0,0,0,0.2);">
<div style="display: flex; justify-content: space-between; align-items: center; border-bottom: 1px solid rgba(255,255,255,0.05); padding-bottom: 12px; margin-bottom: 12px;">
<div>
<span style="color: #38bdf8; font-size: 18px; font-weight: 800; letter-spacing: 0.5px;">#{o.get('reference_id')}</span>
<span style="color: #94a3b8; font-size: 12px; margin-right: 8px;">📅 {o_date}</span>
</div>
<span style="background: {border_color}22; border: 1px solid {border_color}55; color: {border_color}; padding: 4px 12px; border-radius: 20px; font-size: 12px; font-weight: bold;">{status_name}</span>
</div>
<div style="display: flex; flex-wrap: wrap; gap: 10px; font-size: 13px; color: #cbd5e1; margin-bottom: 12px;">
<div style="flex: 1; min-width: 120px;">
<div style="margin-bottom: 6px;">👤 <b style="color:#fff;">العميل:</b> {c_name}</div>
<div style="margin-bottom: 6px;">📍 <b style="color:#fff;">المدينة:</b> {city}</div>
<div>🏢 <b style="color:#fff;">الفرع:</b> {branch}</div>
</div>
<div style="flex: 1; min-width: 120px;">
<div style="margin-bottom: 6px;">💳 <b style="color:#fff;">الدفع:</b> {o.get('payment_method', 'غير محدد')}</div>
<div style="margin-bottom: 6px;">🔗 <b style="color:#fff;">المصدر:</b> {utm_source}</div>
<div>🚚 <b style="color:#fff;">الشحن:</b> {shipping_company}</div>
</div>
</div>
<div style="background: rgba(0,0,0,0.2); border-radius: 8px; padding: 10px; display: flex; justify-content: space-around; text-align: center; border: 1px solid rgba(255,255,255,0.05);">
<div><span style="display:block; font-size:11px; color:#94a3b8;">مجموع السلة</span><b style="color:#fff; font-size:14px;">{round(subtotal, 2):,}</b></div>
<div><span style="display:block; font-size:11px; color:#94a3b8;">الخصومات</span><b style="color:#ef4444; font-size:14px;">{round(total_discount, 2):,}</b></div>
<div><span style="display:block; font-size:11px; color:#94a3b8;">الشحن</span><b style="color:#38bdf8; font-size:14px;">{round(shipping_cost, 2):,}</b></div>
<div><span style="display:block; font-size:11px; color:#94a3b8;">الضريبة</span><b style="color:#eab308; font-size:14px;">{round(tax, 2):,}</b></div>
<div><span style="display:block; font-size:11px; color:#94a3b8;">الإجمالي النهائي</span><b style="color:#22c55e; font-size:15px;">{round(o_total, 2):,}</b></div>
</div>
</div>
"""
                st.markdown(card_html, unsafe_allow_html=True)
                
                with st.expander("🛒 عرض المنتجات"):
                    for item in o.get('items', []):
                        st.markdown(f"- `{item.get('sku', 'بدون SKU')}` | {item.get('name')} (الكمية: **{item.get('quantity', 1)}**)")
